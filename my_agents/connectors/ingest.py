"""Structured + unstructured intake → citable evidence in the project workspace.

Two kinds of input from any channel:

- **structured** — project name, company and free-form fields → the workspace's
  ``project`` / ``business`` context, so the pipeline starts from a stated goal
  instead of an empty brief;
- **unstructured** — uploaded documents → ``artifacts/uploaded_<name>_<ts>.md``
  in the project workspace, recorded in the context as source documents so the
  brief can point the agent at them.

Everything is workspace-scoped: nothing is ever read from or written to another
project. Text-like files decode directly, PDF and DOCX are extracted with the
optional `pypdf` / `python-docx` readers, and anything else is reported as
unsupported rather than silently ingested as garbage.
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from . import config

_STEM_RX = re.compile(r"[^a-zA-Z0-9_-]+")


def _stem(name: str) -> str:
    raw = Path(str(name or "document")).name
    stem = _STEM_RX.sub("_", raw.rsplit(".", 1)[0]).strip("_")
    return (stem or "document")[:60]


def _pdf_text(data: bytes) -> tuple[str, str]:
    """``(text, error)`` for a PDF, via the optional `pypdf` reader."""
    try:
        import io

        from pypdf import PdfReader
    except Exception:
        return "", ("PDF support needs the 'pypdf' package — "
                    "pip install pypdf (see the connector README)")
    try:
        reader = PdfReader(io.BytesIO(data))
        pages: list[str] = []
        for page in reader.pages:
            try:
                pages.append((page.extract_text() or "").strip())
            except Exception:
                pages.append("")
    except Exception as e:
        return "", f"could not read the PDF ({type(e).__name__}) — is it a valid PDF?"
    text = "\n\n".join(p for p in pages if p)
    if not text.strip():
        return "", ("this PDF has no text layer (it looks like a scan) — run OCR or "
                    "export it as text first")
    return text, ""


def _docx_text(data: bytes) -> tuple[str, str]:
    """``(text, error)`` for a .docx, via the optional `python-docx` reader."""
    try:
        import io

        import docx
    except Exception:
        return "", ("docx support needs the 'python-docx' package — "
                    "pip install python-docx (see the connector README)")
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text.strip() for p in document.paragraphs if (p.text or "").strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [(c.text or "").strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    except Exception as e:
        return "", f"could not read the document ({type(e).__name__}) — is it a valid .docx?"
    if not parts:
        return "", "this document has no readable text"
    return "\n".join(parts), ""


# Legacy binary Office formats share a suffix with the modern ones but are a
# different (unreadable) container — say so instead of "unsupported".
_LEGACY = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}


def read_text(name: str, data: bytes, exts: tuple[str, ...] | None = None) -> tuple[str, str]:
    """Return ``(text, error)`` for one upload.

    Text-like files decode directly; PDF and DOCX are extracted with the optional
    readers. Anything else, or a file no text can be got out of, is rejected with a
    reason rather than silently ingested as garbage.
    """
    allowed = exts or config.upload_exts()
    suffix = Path(str(name or "")).suffix.lower()
    if suffix in _LEGACY:
        return "", (f"{suffix} is the old binary format and cannot be read — save it "
                    f"as {_LEGACY[suffix]} or PDF first")
    if suffix not in allowed:
        return "", (f"unsupported file type {suffix or '(none)'} — accepted: "
                    f"{', '.join(allowed)}")
    if len(data) > config.upload_max_bytes():
        return "", f"file too large ({len(data)} bytes)"
    if suffix == ".pdf":
        return _pdf_text(data)
    if suffix == ".docx":
        return _docx_text(data)
    try:
        return data.decode("utf-8", errors="replace"), ""
    except Exception as e:
        return "", f"could not read as text: {type(e).__name__}"


def _context(ws) -> dict:
    ctx = ws.get_context()
    return ctx if isinstance(ctx, dict) else {}


def set_structured(workspace_id: str, fields: dict[str, Any]) -> dict:
    """Write structured intake fields into the workspace project context."""
    from shared.workspace import ProjectWorkspace

    ws = ProjectWorkspace(workspace_id)
    ctx = _context(ws)
    project = ctx.get("project") if isinstance(ctx.get("project"), dict) else {}
    business = ctx.get("business") if isinstance(ctx.get("business"), dict) else {}

    name = str(fields.get("project_name") or fields.get("name") or "").strip()
    goal = str(fields.get("goal") or fields.get("description") or "").strip()
    company = str(fields.get("company") or "").strip()
    if name:
        project["name"] = name[:120]
    if goal:
        project["goal"] = goal[:2000]
    if company:
        business["company"] = company[:200]
    extra = fields.get("extra")
    if isinstance(extra, dict):
        for key, value in list(extra.items())[:40]:
            k = str(key).strip()[:60]
            if k and value not in (None, ""):
                business[k] = str(value)[:500]

    ws.update_context({"project": project, "business": business})
    return {"ok": True, "workspace": ws.project_id,
            "project": project, "business_fields": sorted(business)}


def upload_documents(workspace_id: str, docs: list[tuple[str, str]]) -> dict:
    """Store uploaded document texts in the workspace and record them as sources."""
    from shared.workspace import ProjectWorkspace

    ws = ProjectWorkspace(workspace_id)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    saved: list[dict] = []
    rejected: list[dict] = []
    for name, text in docs:
        if not (text or "").strip():
            rejected.append({"name": str(name), "error": "empty document"})
            continue
        fname = f"uploaded_{_stem(name)}_{ts}.md"
        try:
            ws.save_artifact(fname, text)
        except Exception as e:
            rejected.append({"name": str(name), "error": type(e).__name__})
            continue
        saved.append({"name": str(name), "file": fname, "chars": len(text)})

    if saved:
        ctx = _context(ws)
        business = ctx.get("business") if isinstance(ctx.get("business"), dict) else {}
        docs_ctx = business.get("source_documents")
        if not isinstance(docs_ctx, list):
            docs_ctx = []
        docs_ctx.extend(saved)
        business["source_documents"] = docs_ctx[-50:]
        ws.update_context({"business": business})

    return {"ok": not rejected, "workspace": ws.project_id,
            "files": [s["file"] for s in saved], "sources": saved,
            "rejected": rejected}


def ingest_files(workspace_id: str, files: list[tuple[str, bytes]]) -> dict:
    """Decode raw uploads, then store the readable ones."""
    texts: list[tuple[str, str]] = []
    rejected: list[dict] = []
    for name, data in files:
        text, err = read_text(name, data)
        if err:
            rejected.append({"name": str(name), "error": err})
        else:
            texts.append((name, text))
    out = upload_documents(workspace_id, texts)
    if rejected:
        out["rejected"] = list(out.get("rejected") or []) + rejected
        out["ok"] = False
    return out


def set_mode(workspace_id: str, mode: str = "generation") -> dict:
    """Seed the conversation mode in the workspace context.

    ``generation`` makes the pipeline start producing documents on the first
    turn instead of opening a discovery chat. The orchestrator hydrates this
    context when it binds the workspace, so the intake form can go straight to
    work — which also means the run needs no discovery model call.
    """
    from shared.workspace import ProjectWorkspace

    want = (mode or "").strip().lower()
    if want not in ("generation", "discovery"):
        return {"ok": False, "error": "mode must be generation or discovery"}
    ws = ProjectWorkspace(workspace_id)
    ctx = _context(ws)
    conv = ctx.get("conversation") if isinstance(ctx.get("conversation"), dict) else {}
    conv["mode"] = want
    ws.update_context({"conversation": conv})
    return {"ok": True, "workspace": ws.project_id, "mode": want}


def source_documents(workspace_id: str) -> list[dict]:
    """The source documents recorded for a workspace (for the web page/tests)."""
    from shared.workspace import ProjectWorkspace

    ws = ProjectWorkspace(workspace_id)
    business = _context(ws).get("business")
    if isinstance(business, dict) and isinstance(business.get("source_documents"), list):
        return [d for d in business["source_documents"] if isinstance(d, dict)]
    return []
